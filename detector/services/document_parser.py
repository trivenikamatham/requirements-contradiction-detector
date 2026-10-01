import io
import re
from typing import List, Dict, Any, Optional

try:
    import pypdf
except ImportError:
    try:
        import PyPDF2 as pypdf
    except ImportError:
        pypdf = None

try:
    import docx
except ImportError:
    docx = None


class DocumentParserError(Exception):
    """Exception raised for errors during document parsing."""
    pass


class DocumentParser:
    """
    Extracts text from PDF, DOCX, and TXT requirements documents
    and splits content into individual requirements while preserving
    provenance (document name, page number, and section references).
    """

    # Common requirement prefixes (e.g., REQ-001, FR-1.2, BR-3, NFR-05, etc.)
    REQ_ID_PATTERN = re.compile(
        r'^\s*(?:\[?\s*(?:REQ|FR|BR|NFR|R|SR|UR|SYS|FUNC|SEC)[-_]?[A-Z0-9.\-]+\s*\]?'
        r'|(?:\d+\.){1,4}\d+'
        r'|\d+\.?\s*(?:\)|\.)'
        r'|[•\*\-\–\—])\s*',
        re.IGNORECASE
    )

    # Specific pattern to extract explicit identifier if present
    EXPLICIT_ID_PATTERN = re.compile(
        r'^\s*\[?\s*((?:REQ|FR|BR|NFR|R|SR|UR|SYS|FUNC|SEC)[-_]?[A-Za-z0-9.\-]+)\s*\]?[:\.\s\-]+',
        re.IGNORECASE
    )

    # Keywords typical in IEEE 830 / RFC 2119 requirement statements
    REQUIREMENT_KEYWORDS = [
        'shall', 'must', 'will', 'should', 'requires', 'required', 'is required',
        'mandatory', 'shall not', 'must not', 'prohibited', 'cannot', 'is prohibited'
    ]

    # Boilerplate lines to ignore
    BOILERPLATE_PATTERNS = [
        re.compile(r'^(?:table of contents|contents|index|appendix)\b', re.IGNORECASE),
        re.compile(r'^(?:page \d+|page \d+ of \d+|\d+\s*/\s*\d+)$', re.IGNORECASE),
        re.compile(r'^(?:confidential|proprietary|all rights reserved)\b', re.IGNORECASE),
        re.compile(r'^(?:document control|version history|revision history)\b', re.IGNORECASE),
    ]

    def parse_file(self, file_path_or_obj, filename: str) -> List[Dict[str, Any]]:
        """
        Parses a file object or path and returns a list of requirement dictionaries.
        Each dictionary contains:
            - 'req_id': identifier string (e.g. 'REQ-01' or 'BR-04')
            - 'text': full statement
            - 'page_number': int or None
            - 'section': str or None
        """
        ext = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''
        
        if ext == 'pdf':
            return self._parse_pdf(file_path_or_obj, filename)
        elif ext == 'docx':
            return self._parse_docx(file_path_or_obj, filename)
        elif ext in ('txt', 'md'):
            return self._parse_txt(file_path_or_obj, filename)
        else:
            raise DocumentParserError(f"Unsupported file format '.{ext}'. Supported formats: .pdf, .docx, .txt")

    def _parse_pdf(self, file_source, filename: str) -> List[Dict[str, Any]]:
        """Extracts text page by page from PDF and isolates requirements."""
        if pypdf is None:
            raise DocumentParserError("PDF extraction library (pypdf) is not available.")

        try:
            if hasattr(file_source, 'read'):
                file_source.seek(0)
                reader = pypdf.PdfReader(file_source)
            else:
                reader = pypdf.PdfReader(file_source)

            requirements = []
            req_counter = 1
            current_section = "General Requirements"

            for page_idx, page in enumerate(reader.pages):
                page_num = page_idx + 1
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""

                lines = [line.strip() for line in text.split('\n') if line.strip()]
                page_reqs, current_section, req_counter = self._extract_from_lines(
                    lines, page_num=page_num, default_section=current_section, counter_start=req_counter
                )
                requirements.extend(page_reqs)

            # If no requirements were extracted with strict heuristics, fall back to paragraph chunking
            if not requirements:
                requirements = self._fallback_chunking(reader, filename)

            return requirements
        except Exception as e:
            raise DocumentParserError(f"Failed to parse PDF '{filename}': {str(e)}")

    def _parse_docx(self, file_source, filename: str) -> List[Dict[str, Any]]:
        """Extracts paragraphs and headings from DOCX file."""
        if docx is None:
            raise DocumentParserError("DOCX extraction library (python-docx) is not available.")

        try:
            if hasattr(file_source, 'read'):
                file_source.seek(0)
                doc = docx.Document(file_source)
            else:
                doc = docx.Document(file_source)

            requirements = []
            current_section = "Overview"
            req_counter = 1
            page_estimate = 1
            line_count = 0

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                line_count += 1
                if line_count % 35 == 0:
                    page_estimate += 1

                # Check if paragraph is a heading
                style_name = para.style.name.lower() if para.style else ""
                if 'heading' in style_name or self._is_section_header(text):
                    current_section = text
                    continue

                if self._is_boilerplate(text):
                    continue

                req_item = self._extract_requirement_from_text(
                    text, page_num=page_estimate, section=current_section, counter=req_counter
                )
                if req_item:
                    requirements.append(req_item)
                    req_counter += 1

            # Also inspect tables if any exist in the docx
            for table_idx, table in enumerate(doc.tables):
                for row_idx, row in enumerate(table.rows):
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    row_combined = " | ".join(row_texts)
                    if self._is_boilerplate(row_combined) or len(row_combined) < 20:
                        continue
                    if any(k in row_combined.lower() for k in self.REQUIREMENT_KEYWORDS):
                        req_item = self._extract_requirement_from_text(
                            row_combined,
                            page_num=page_estimate,
                            section=f"Table {table_idx + 1} (Row {row_idx + 1})",
                            counter=req_counter
                        )
                        if req_item:
                            requirements.append(req_item)
                            req_counter += 1

            return requirements
        except Exception as e:
            raise DocumentParserError(f"Failed to parse DOCX '{filename}': {str(e)}")

    def _parse_txt(self, file_source, filename: str) -> List[Dict[str, Any]]:
        """Extracts requirements from plain text or markdown files."""
        try:
            if hasattr(file_source, 'read'):
                file_source.seek(0)
                raw_bytes = file_source.read()
                if isinstance(raw_bytes, str):
                    content = raw_bytes
                else:
                    try:
                        content = raw_bytes.decode('utf-8')
                    except UnicodeDecodeError:
                        content = raw_bytes.decode('latin-1', errors='ignore')
            elif isinstance(file_source, str):
                import os
                if os.path.exists(file_source):
                    with open(file_source, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()
                else:
                    content = file_source
            else:
                content = str(file_source)

            lines = [l.strip() for l in content.splitlines()]
            requirements, _, _ = self._extract_from_lines(
                lines, page_num=1, default_section="Requirements Specification", counter_start=1
            )
            return requirements
        except Exception as e:
            raise DocumentParserError(f"Failed to parse text document '{filename}': {str(e)}")

    def _extract_from_lines(self, lines: List[str], page_num: Optional[int], default_section: str, counter_start: int):
        """Processes lines into structured requirement entries."""
        requirements = []
        current_section = default_section
        counter = counter_start
        current_req_text = ""
        current_req_id = None

        def flush_current():
            nonlocal current_req_text, current_req_id, counter
            if current_req_text:
                clean_text = self._clean_requirement_text(current_req_text)
                if len(clean_text) >= 15:
                    requirements.append({
                        'req_id': current_req_id or f"REQ-{counter:03d}",
                        'text': clean_text,
                        'page_number': page_num,
                        'section': current_section
                    })
                    counter += 1
            current_req_text = ""
            current_req_id = None

        for line in lines:
            if not line:
                flush_current()
                continue

            if self._is_boilerplate(line):
                continue

            # Check if this line is a section heading
            if self._is_section_header(line):
                flush_current()
                current_section = re.sub(r'^[#=\-\*\s\d\.:]+', '', line).strip() or line
                continue

            # Check for explicit ID match (e.g., REQ-01:, FR-2.1 - The system shall...)
            id_match = self.EXPLICIT_ID_PATTERN.match(line)
            if id_match:
                flush_current()
                current_req_id = id_match.group(1).upper()
                current_req_text = line[id_match.end():].strip()
                if not current_req_text:
                    current_req_text = line
                continue

            # Check for bullet/number prefix or requirement statement starter
            is_prefixed = bool(self.REQ_ID_PATTERN.match(line))
            has_keywords = any(kw in line.lower() for kw in self.REQUIREMENT_KEYWORDS)

            if is_prefixed or (has_keywords and not current_req_text):
                flush_current()
                # Clean prefix for text body
                stripped = self.REQ_ID_PATTERN.sub('', line).strip()
                current_req_text = stripped if stripped else line
                # Derive an ID if prefix had numbers (e.g. 3.1.2)
                num_match = re.match(r'^\s*(\d+(?:\.\d+)+|\d+[\.\)])\s*', line)
                if num_match:
                    num_id = num_match.group(1).rstrip('.)')
                    current_req_id = f"SEC-{num_id}"
                else:
                    current_req_id = f"REQ-{counter:03d}"
            elif current_req_text:
                # Continuation of current requirement statement
                current_req_text += " " + line
            elif has_keywords:
                flush_current()
                current_req_text = line
                current_req_id = f"REQ-{counter:03d}"

        flush_current()
        return requirements, current_section, counter

    def _extract_requirement_from_text(self, text: str, page_num: Optional[int], section: str, counter: int) -> Optional[Dict[str, Any]]:
        """Extracts an individual requirement item from a single paragraph."""
        id_match = self.EXPLICIT_ID_PATTERN.match(text)
        if id_match:
            req_id = id_match.group(1).upper()
            body = text[id_match.end():].strip() or text
            return {
                'req_id': req_id,
                'text': self._clean_requirement_text(body),
                'page_number': page_num,
                'section': section
            }

        is_prefixed = bool(self.REQ_ID_PATTERN.match(text))
        has_keywords = any(kw in text.lower() for kw in self.REQUIREMENT_KEYWORDS)

        if is_prefixed or has_keywords:
            cleaned = self.REQ_ID_PATTERN.sub('', text).strip() or text
            if len(cleaned) >= 15:
                num_match = re.match(r'^\s*(\d+(?:\.\d+)+|\d+[\.\)])\s*', text)
                if num_match:
                    req_id = f"SEC-{num_match.group(1).rstrip('.)')}"
                else:
                    req_id = f"REQ-{counter:03d}"
                return {
                    'req_id': req_id,
                    'text': self._clean_requirement_text(cleaned),
                    'page_number': page_num,
                    'section': section
                }
        return None

    def _is_section_header(self, line: str) -> bool:
        """Determines if a line represents a section heading."""
        if line.startswith(('#', '==', '--', 'Section', 'SECTION', 'Chapter', 'CHAPTER')):
            return True
        if re.match(r'^\d+(?:\.\d+){0,3}\s+[A-Z][a-zA-Z\s]{2,40}$', line):
            return True
        if len(line) < 60 and line.isupper() and not any(kw in line.lower() for kw in self.REQUIREMENT_KEYWORDS):
            return True
        return False

    def _is_boilerplate(self, line: str) -> bool:
        """Detects noise, footers, headers, or metadata lines."""
        for pattern in self.BOILERPLATE_PATTERNS:
            if pattern.search(line):
                return True
        return False

    def _clean_requirement_text(self, text: str) -> str:
        """Removes excess whitespace, leading bullet marks, and trailing linefeeds."""
        text = re.sub(r'\s+', ' ', text).strip()
        text = re.sub(r'^[•\*\-\–\—\s:]+', '', text).strip()
        return text

    def _fallback_chunking(self, reader, filename: str) -> List[Dict[str, Any]]:
        """Fallback to sentence/paragraph chunking if strict pattern matching yields nothing."""
        requirements = []
        counter = 1
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            # Split into paragraphs
            paras = [p.strip() for p in page_text.split('\n\n') if len(p.strip()) > 25]
            for para in paras:
                if not self._is_boilerplate(para):
                    requirements.append({
                        'req_id': f"REQ-{counter:03d}",
                        'text': self._clean_requirement_text(para),
                        'page_number': page_idx + 1,
                        'section': "General Requirements"
                    })
                    counter += 1
        return requirements
